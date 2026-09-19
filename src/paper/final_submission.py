"""Final Applied Sciences submission artifacts built from the frozen P16 manuscript (no scientific change).

Sources of truth: `paper/manuscript/manuscript_p16_final.md` (text; rendered in `paper/submission_p16/`) and the P16
package (tables, figures, supplementary files). This module only formats them:
- the main manuscript DOCX in the Applied Sciences Word template (the template is supplied at run time and never
  copied into the repository: the journal restricts it to submission use);
- a supplementary DOCX (Tables S1-S44 list with captions, the Table S19 reproduction record, Figures S1-S4) and a
  workbook with every supplementary table at full precision (8,513 rows, up to 37 columns: not legible as Word
  pages, and not reduced);
- Figures 1-6 as separate files, byte-identical to the P16 package;
- a manifest with the SHA-256 of every final artifact.
Structural, privacy and scientific-equivalence checks read the DOCX files. Page rendering for visual QA uses Microsoft
Word (scripts/render_docx_pages.ps1) and is recorded in the QA reports under `paper/submission_final/qa/`.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from src.data import paths
from src.data.io_guard import open_for_write, read_bytes, write_json
from src.paper import docx_export as DX
from src.paper import p16_submission as P16

ROOT = paths.PROJECT_ROOT
FINAL = ROOT / "paper" / "submission_final"
SOURCE_COMMIT = "62b1618c5df5726d1bc3b9cf7277e7dc7db35a2f"
SOURCE_MANUSCRIPT = "paper/manuscript/manuscript_p16_final.md"
P16_PACKAGE = P16.PACKAGE
P16_RENDERED = P16.RENDERED
MAIN_DOCX = FINAL / "manuscript" / "Applied_Sciences_SmartMat_Final.docx"
SUPP_DOCX = FINAL / "supplementary" / "Supplementary_Materials.docx"
SUPP_XLSX = FINAL / "supplementary" / "Supplementary_Tables_S1-S44.xlsx"
FIGURES = FINAL / "figures"
COVER_LETTER = FINAL / "cover_letter" / "cover_letter_final_draft.md"
METADATA = FINAL / "metadata" / "FINAL_METADATA_CHECKLIST.md"
REFERENCE_CHECK = FINAL / "metadata" / "REFERENCE_FINAL_CHECK.md"
EXTERNAL_REQUEST = FINAL / "metadata" / "EXTERNAL_INFORMATION_REQUEST.md"
CHECKLIST = FINAL / "SUBMISSION_CHECKLIST.md"
QA_REPORTS = {name: FINAL / "qa" / name for name in ("DOCX_STRUCTURAL_QA.md", "DOCX_VISUAL_QA.md",
                                                     "SUPPLEMENTARY_QA.md")}
MANIFEST = FINAL / "MANIFEST.json"
N_TABLES, N_FIGURES, N_SUPP_TABLES, N_SUPP_FIGURES = 11, 6, 44, 4
STATUS = "submission_artifacts_complete_metadata_pending"
PLACEHOLDER = re.compile(r"\[CONFIRM BEFORE SUBMISSION: ([^\]]+)\]")
PRIVATE = {
    "local path": re.compile(r"[A-Za-z]:\\\\|[A-Za-z]:\\|[A-Za-z]:/Users|/Users/|/home/|\\\\Users\\\\"),
    "raw package or journal folder": re.compile(r"스마트|학술지|Desktop"),
    "e-mail address": re.compile(r"[\w.+-]+@[\w-]+\.[A-Za-z]{2,}"),
    "calendar date": re.compile(r"\b(19|20)\d{2}-\d{2}-\d{2}\b"),
}
SENSOR_TERMS = re.compile(r"physiological microclimate|body[- ]bed interface|skin[- ]interface", re.I)
NOT_ALLOWED_ANYWHERE = re.compile(r"available (up)?on request|ready[_ ]to[_ ]submit", re.I)
COVER_HYPE = re.compile(r"\bfirst\b|state[- ]of[- ]the[- ]art|superior|\bproves?\b|clinical deployment|"
                        r"negative transfer", re.I)
OPEN_STATUS = "OPEN — EXTERNAL CONFIRMATION REQUIRED"
N_METADATA_ITEMS = 18
P8_COVER_LETTER_SHA256 = "552aaa47e5a27cdb6da14651e711922533fda0181eb0e3f38d9b980690041d5b"   # kept unchanged
FORBIDDEN_TEXT = re.compile(r"Figure S5|figureS5|night_series|night_events|\bTODO\b|\bFIXME\b|mock review|"
                            r"internal review|\bP1[0-9]\b", re.I)


def title() -> str:
    return P16.FINAL_TITLE


# --- main manuscript --------------------------------------------------------------------------------------------

def rendered_markdown() -> str:
    return P16_RENDERED.read_text(encoding="utf-8")


def build_main(template: Path, out: Path = MAIN_DOCX) -> Path:
    return DX.build_docx(template, out, markdown=rendered_markdown(), figures_dir=P16_PACKAGE / "figures",
                         layout=True)


def main_figures() -> list[Path]:
    return sorted((P16_PACKAGE / "figures").glob("figure[1-6]_*.png"))


def copy_figures(out_dir: Path = FIGURES) -> list[Path]:
    """Figures 1-6 as separate submission files, byte-identical to the P16 package."""
    out = []
    for src in main_figures():
        with open_for_write(out_dir / src.name, "wb") as fh:
            fh.write(read_bytes(src))
        out.append(out_dir / src.name)
    return out


# --- supplementary ----------------------------------------------------------------------------------------------

def supp_table_index() -> list[tuple[str, str, list[str]]]:
    """(Table Sn, title, [file names]) from the P16 supplementary index."""
    out = []
    for line in (P16_PACKAGE / "supplementary" / "README.md").read_text(encoding="utf-8").splitlines():
        m = re.match(r"\| Table (S\d+) \| (.+?) \| (.+) \|$", line)
        if m:
            out.append((m.group(1), m.group(2), re.findall(r"`([^`]+)`", m.group(3))))
    return out


def supp_figure_captions() -> dict[str, str]:
    text = " ".join(rendered_markdown().split("## Supplementary Materials", 1)[1].split("\n## ", 1)[0].split())
    return {m.group(1): m.group(2).strip() for m in
            re.finditer(r"Figure (S\d+): (.+?); (?=Figure S\d+:|Table S1:)", text)}


def sheet_name(stem: str, used: set[str]) -> str:
    base = re.sub(r"^TableS", "S", stem)[:31]
    name, k = base, 2
    while name in used:
        name = f"{base[:28]}_{k}"
        k += 1
    used.add(name)
    return name


def sheet_plan() -> list[tuple[str, str, str, Path]]:
    """(table id, title, sheet name, CSV path) for every supplementary CSV, in table order."""
    used: set[str] = {"Index"}
    plan = []
    for sid, ttl, files in supp_table_index():
        for f in files:
            if f.endswith(".csv"):
                plan.append((sid, ttl, sheet_name(Path(f).stem, used),
                             P16_PACKAGE / "supplementary" / "tables" / f))
    return plan


def _cell(value: str):
    """Exactly the CSV text (spreadsheet floats keep only about 15 significant digits, which would not be lossless)."""
    return None if value == "" else value


def build_workbook(out: Path = SUPP_XLSX) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    wb.properties.creator = ""
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    index = wb.active
    index.title = "Index"
    index.append([f"Supplementary Tables S1–S44 — {title()}"])
    index.append(["Every table at full precision: each value is stored as text exactly as in the source table. "
                  "No calendar dates. Table S19 is in the supplementary document."])
    index.append([])
    index.append(["Table", "Content", "Worksheet", "Source file"])
    for sid, ttl, sheet, src in sheet_plan():
        index.append([sid, ttl, sheet, src.name])
        ws = wb.create_sheet(sheet)
        with open(src, encoding="utf-8", newline="") as fh:
            for row in csv.reader(fh):
                ws.append([_cell(v) for v in row])
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as src, open_for_write(out, "wb") as fh:    # fixed entry times: deterministic bytes
        with zipfile.ZipFile(fh, "w", zipfile.ZIP_DEFLATED) as z:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "docProps/core.xml":                  # openpyxl stamps the save time
                    data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
                z.writestr(zipfile.ZipInfo(info.filename, DX.FIXED_TIME), data, compress_type=zipfile.ZIP_DEFLATED)
    return out


def supplementary_markdown() -> str:
    caps = supp_figure_captions()
    lines = ["# Supplementary Materials", "", f"**Manuscript:** {title()}", "",
             f"**Note:** Tables S1–S44 are provided in full, at full precision, in the accompanying workbook "
             f"{SUPP_XLSX.name} (one worksheet per table part; the worksheet names are listed below). Table S19, "
             "a reproduction record, is reproduced in this document. Figures S1–S4 follow. No calendar date is "
             "included.", "", "## Supplementary Tables", ""]
    sheets: dict[str, list[str]] = {}
    for sid, _, sheet, _ in sheet_plan():
        sheets.setdefault(sid, []).append(sheet)
    for sid, ttl, files in supp_table_index():
        where = ("this document" if sid == "S19" else "worksheets " + ", ".join(sheets.get(sid, [])))
        lines.append(f"- **Table {sid}.** {ttl} ({where}).")
    lines += ["", "## Reproduction Record (Table S19)", ""]
    s19 = (P16_PACKAGE / "supplementary" / "tables" / "TableS19_reproduction_record.md").read_text(encoding="utf-8")
    for line in s19.splitlines():
        if line.startswith("# "):
            continue
        lines.append("#### " + line.lstrip("#").strip() if line.startswith("#") else line)
    lines += ["", "## Supplementary Figures", ""]
    for n in range(1, N_SUPP_FIGURES + 1):
        stem = next(p.stem for p in (P16_PACKAGE / "supplementary" / "figures").glob(f"figureS{n}_*.png"))
        lines += [f"![Figure S{n}](../figures/{stem}.png)", "", f"**Figure S{n}.** {caps[f'S{n}']}", ""]
    return "\n".join(lines) + "\n"


def build_supplementary(template: Path, out: Path = SUPP_DOCX) -> Path:
    return DX.build_docx(template, out, markdown=supplementary_markdown(),
                         figures_dir=P16_PACKAGE / "supplementary" / "figures", layout=True,
                         article_type=None, ragged=True)


# --- checks on the DOCX files -------------------------------------------------------------------------------------

def _parts(docx: Path) -> dict[str, str]:
    z = zipfile.ZipFile(docx)
    return {n: z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels"))}


def docx_text(docx: Path) -> str:
    return DX.docx_text(docx)


def structure(docx: Path) -> dict:
    doc = _parts(docx)["word/document.xml"]
    text = docx_text(docx)
    return {"tables": doc.count("<w:tbl>"), "inline_images": doc.count("<wp:inline"),
            "table_captions": re.findall(r"^Table (S?\d+)\. ", text, re.M),
            "figure_captions": re.findall(r"^Figure (S?\d+)\. ", text, re.M),
            "headings": [re.sub(r"<[^>]+>", "", h).replace("&amp;", "&") for h in re.findall(
                r'<w:pStyle w:val="MDPI2[123]heading[123]"/>(?:(?!</w:pPr>).)*</w:pPr>(.*?)</w:p>', doc)],
            "references": doc.count('w:val="MDPI81references"'),
            "placeholders": [" ".join(m.group(1).split()) for m in PLACEHOLDER.finditer(text)],
            "paragraphs": doc.count("<w:p>") + doc.count("<w:p "),
            "tracked_changes": len(re.findall(r"<w:(ins|del) ", doc)),
            "comments_part": any(n.startswith("word/comments") for n in _parts(docx)),
            "hidden_text": doc.count("<w:vanish/>")}


def privacy_problems(docx: Path) -> list[str]:
    out = []
    for name, xml in _parts(docx).items():
        text = re.sub(r"<[^>]+>", " ", xml) if name.endswith(".xml") else xml
        scan = text if name != "word/document.xml" else text.split("References", 1)[0]
        for label, rx in PRIVATE.items():
            for m in rx.finditer(scan):
                if label == "e-mail address" and m.group(0).endswith(("mdpi.com", "w3.org", "openxmlformats.org",
                                                                      "microsoft.com", "purl.org")):
                    continue
                out.append(f"{docx.name}:{name}: {label}: {m.group(0)!r}")
        if re.search(r"<w:(ins|del) ", xml):
            out.append(f"{docx.name}:{name}: tracked changes")
        if "<w:vanish/>" in xml:
            out.append(f"{docx.name}:{name}: hidden text")
        if "attachedTemplate" in xml:
            out.append(f"{docx.name}:{name}: attached template path")
    if any(n.startswith("word/comments") for n in _parts(docx)):
        out.append(f"{docx.name}: comments part present")
    core = _parts(docx).get("docProps/core.xml", "")
    out += [f"{docx.name}: document property {m}" for m in _person_properties(core)]
    return out


def _person_properties(core_xml: str) -> list[str]:
    """Author-like document properties with a value (none may be generated before the authors are confirmed)."""
    return [m.group(1) for m in re.finditer(r"<(?:dc|cp):(creator|lastModifiedBy)>([^<]+)</", core_xml)
            if m.group(2).strip()]


def workbook_privacy_problems(xlsx: Path = SUPP_XLSX) -> list[str]:
    z = zipfile.ZipFile(xlsx)
    out = [f"{xlsx.name}: document property {m}" for m in _person_properties(
        z.read("docProps/core.xml").decode("utf-8"))]
    for name in z.namelist():
        text = z.read(name).decode("utf-8", "replace")
        out += [f"{xlsx.name}:{name}: {label}" for label, rx in PRIVATE.items()
                if label in ("local path", "raw package or journal folder") and rx.search(text)]
    return out


def _numbers(text: str) -> Counter:
    return Counter(re.findall(r"[+−-]?\d[\d,]*(?:\.\d+)?", text))


def _plain(markdown: str) -> str:
    md = re.sub(r"<!--.*?-->", "", markdown, flags=re.S)
    md = re.sub(r"!\[[^\]]*\]\([^)]*\)", " ", md)
    md = re.sub(r"(?m)^\s*\d+\.\s+(?=\S)", " ", md)       # list markers (Word numbers lists itself)
    md = re.sub(r"[*`#|]", " ", md)
    return md.replace("\\|", "|")


def equivalence_problems(docx: Path = MAIN_DOCX) -> list[str]:
    """The DOCX carries the P16 text: same title, abstract, conclusions and the same multiset of numbers (line
    numbers and page furniture live in the template, not in the body)."""
    md = rendered_markdown()
    body_md = md.split("## References")[0]
    text = docx_text(docx)
    body_docx = text.split("\nReferences\n")[0] if "\nReferences\n" in text else text
    out = []
    flat_docx = " ".join(body_docx.split())
    if title() not in flat_docx:
        out.append("final title missing from the DOCX")
    abstract = " ".join(body_md.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0].split())
    if abstract not in flat_docx:
        out.append("abstract text differs from P16")
    keywords = " ".join(body_md.split("**Keywords:**", 1)[1].split("\n## ", 1)[0].split())
    if f"Keywords: {keywords}" not in flat_docx:
        out.append("keywords differ from P16")
    concl = " ".join(_plain(body_md.split("## 6. Conclusions", 1)[1].split("## Supplementary Materials")[0]).split())
    if concl not in " ".join(_plain(body_docx).split()):
        out.append("conclusions text differs from P16")
    n_md = _numbers(_plain(body_md.replace("\n## Abstract", "\nAbstract: ")))
    n_docx = _numbers(body_docx)
    for k in set(n_md) | set(n_docx):
        if n_md[k] != n_docx[k]:
            out.append(f"number {k!r}: P16 {n_md[k]}, DOCX {n_docx[k]}")
    return out


def source_headings() -> list[str]:
    """P16 section headings in order, as the template shows them (Abstract and back matter are run-in labels)."""
    skip = {"Abstract", *DX.BACK_MATTER}
    return [h for h in re.findall(r"^#{2,4} (.+)$", rendered_markdown(), re.M) if h.strip() not in skip]


def main_structure_problems(docx: Path = MAIN_DOCX) -> list[str]:
    s = structure(docx)
    out = []
    if s["headings"] != source_headings():
        out.append("heading order or text differs from P16")
    if s["tables"] != N_TABLES + 1:                          # the numbered tables and the abbreviations table
        out.append(f"{s['tables']} tables (expected {N_TABLES} numbered tables and the abbreviations table)")
    if s["table_captions"] != [str(k) for k in range(1, N_TABLES + 1)]:
        out.append(f"table captions {s['table_captions']}")
    if s["inline_images"] != N_FIGURES:
        out.append(f"{s['inline_images']} figures (expected {N_FIGURES})")
    if s["figure_captions"] != [str(k) for k in range(1, N_FIGURES + 1)]:
        out.append(f"figure captions {s['figure_captions']}")
    if s["references"] != 31:
        out.append(f"{s['references']} reference entries (expected 31)")
    text = docx_text(docx)
    out += [f"forbidden text: {m.group(0)!r}" for m in FORBIDDEN_TEXT.finditer(text)]
    out += [f"not allowed: {m.group(0)!r}" for m in NOT_ALLOWED_ANYWHERE.finditer(text)]
    source_terms = Counter(m.group(0).lower() for m in SENSOR_TERMS.finditer(rendered_markdown()))
    docx_terms = Counter(m.group(0).lower() for m in SENSOR_TERMS.finditer(text))
    out += [f"sensor-placement wording not in P16: {t!r}" for t in docx_terms if docx_terms[t] > source_terms[t]]
    expected = [t for _, t in P16.B.placeholders(P16.read_source())]
    if s["placeholders"] != expected:
        out.append("placeholders differ from the P16 source")
    for sec in ("Author Contributions", "Funding", "Institutional Review Board Statement",
                "Informed Consent Statement", "Data Availability Statement", "Acknowledgments",
                "Conflicts of Interest"):
        if f"{sec}:" not in text:
            out.append(f"back matter missing: {sec}")
    if s["tracked_changes"] or s["comments_part"] or s["hidden_text"]:
        out.append("tracked changes, comments or hidden text present")
    return out


def supplementary_problems(docx: Path = SUPP_DOCX, xlsx: Path = SUPP_XLSX) -> list[str]:
    out = []
    s = structure(docx)
    if s["inline_images"] != N_SUPP_FIGURES or s["figure_captions"] != [f"S{k}" for k in range(1, 5)]:
        out.append(f"supplementary figures {s['inline_images']} / captions {s['figure_captions']}")
    text = docx_text(docx)
    listed = re.findall(r"^Table (S\d+)\. ", text, re.M)
    if listed != [f"S{k}" for k in range(1, N_SUPP_TABLES + 1)]:
        out.append(f"supplementary table list {listed}")
    out += [f"forbidden text: {m.group(0)!r}" for m in re.finditer(r"Figure S5|figureS5|night_series|night_events",
                                                                    text)]
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True)
    for sid, _, sheet, src in sheet_plan():
        if sheet not in wb.sheetnames:
            out.append(f"worksheet missing: {sheet}")
            continue
        with open(src, encoding="utf-8", newline="") as fh:
            rows = list(csv.reader(fh))
        got = [["" if v is None else v for v in r] for r in wb[sheet].iter_rows(values_only=True)]
        want = [[_cell(v) if _cell(v) is not None else "" for v in r] for r in rows]
        if got != want:
            out.append(f"worksheet {sheet} differs from {src.name}")
    return out


def figure_problems() -> list[str]:
    srcs = main_figures()
    out = [] if len(srcs) == N_FIGURES else [f"{len(srcs)} P16 main figures (expected {N_FIGURES})"]
    for src in srcs:
        dst = FIGURES / src.name
        if not dst.is_file() or read_bytes(dst) != read_bytes(src):
            out.append(f"figure file missing or not byte-identical: {dst.name}")
    extra = sorted({q.name for q in FIGURES.glob("*")} - {q.name for q in srcs}) if FIGURES.is_dir() else []
    return out + [f"unexpected file in figures/: {n}" for n in extra]


# --- documents written by the authors' tooling ---------------------------------------------------------------------

def cover_letter_problems(path: Path = COVER_LETTER) -> list[str]:
    if not path.is_file():
        return ["cover letter missing"]
    text = path.read_text(encoding="utf-8")
    flat = " ".join(text.split())
    out = [] if title() in flat else ["final title missing from the cover letter"]
    if "Smart-Mat Microclimate Estimation" in flat:
        out.append("old title in the cover letter")
    letter = text.split("\n---\n", 1)[-1]                   # the letter itself, after the status note
    out += [f"overstated wording: {m.group(0)!r}" for m in COVER_HYPE.finditer(letter)]
    out += [f"not allowed: {m.group(0)!r}" for m in NOT_ALLOWED_ANYWHERE.finditer(text)]
    if not PLACEHOLDER.search(letter):
        out.append("external items are not left as placeholders")
    if sha256(ROOT / "paper" / "manuscript" / "COVER_LETTER_DRAFT.md") != P8_COVER_LETTER_SHA256:
        out.append("the P8-era cover letter was modified")
    return out


def metadata_problems(path: Path = METADATA) -> list[str]:
    if not path.is_file():
        return ["metadata checklist missing"]
    text = path.read_text(encoding="utf-8")
    rows = re.findall(r"^\| (M\d\d) \|(.*)\|$", text, re.M)
    out = [] if [r[0] for r in rows] == [f"M{k:02d}" for k in range(1, N_METADATA_ITEMS + 1)] else \
        [f"checklist items {[r[0] for r in rows]}"]
    out += [f"{rid}: status is not open" for rid, rest in rows if OPEN_STATUS not in rest]
    flat = " ".join(text.split())
    out += [f"placeholder not in the checklist: {t}" for _, t in P16.B.placeholders(P16.read_source())
            if f"[CONFIRM BEFORE SUBMISSION: {t}]" not in flat]
    return out


def external_request_problems(path: Path = EXTERNAL_REQUEST) -> list[str]:
    """The hand-off request: every row open, every checklist item M01-M18 and cover-letter item L1-L4 requested."""
    if not path.is_file():
        return ["external information request missing"]
    rows = re.findall(r"^\| ([A-D]\d+) \|.*\| ((?:M|L)\d+) \| (\w+) \|$", path.read_text(encoding="utf-8"), re.M)
    out = [f"{r[0]}: status {r[2]}" for r in rows if r[2] != "OPEN"]
    ids = {r[1] for r in rows}
    want = {f"M{k:02d}" for k in range(1, N_METADATA_ITEMS + 1)} | {f"L{k}" for k in range(1, 5)}
    return out + [f"not requested: {i}" for i in sorted(want - ids)]


def qa_problems(pages: dict[str, int]) -> list[str]:
    """Every rendered page has a row in the visual QA reports, and no report claims submission readiness."""
    out = [f"QA report missing: {p.name}" for p in QA_REPORTS.values() if not p.is_file()]
    if out:
        return out
    for key, report in (("main", QA_REPORTS["DOCX_VISUAL_QA.md"]),
                        ("supplementary", QA_REPORTS["SUPPLEMENTARY_QA.md"])):
        rows = [int(n) for n in re.findall(r"^\| (\d+) \| ", report.read_text(encoding="utf-8"), re.M)]
        if rows != list(range(1, pages.get(key, 0) + 1)):
            out.append(f"{report.name}: page rows {rows[:3]}… do not cover pages 1–{pages.get(key)}")
    for p in [*QA_REPORTS.values(), CHECKLIST]:
        if p.is_file():
            out += [f"{p.name}: not allowed: {m.group(0)!r}" for m in NOT_ALLOWED_ANYWHERE.finditer(
                p.read_text(encoding="utf-8"))]
    return out


# --- manifest -----------------------------------------------------------------------------------------------------

def artifacts() -> list[Path]:
    return [MAIN_DOCX, SUPP_DOCX, SUPP_XLSX, *(FIGURES / f.name for f in main_figures()), COVER_LETTER, METADATA,
            EXTERNAL_REQUEST, REFERENCE_CHECK, CHECKLIST, *QA_REPORTS.values()]


def sha256(p: Path) -> str:
    data = read_bytes(p)
    if p.suffix in (".md", ".json", ".csv"):
        data = data.replace(b"\r\n", b"\n")
    return hashlib.sha256(data).hexdigest()


def write_manifest(pages: dict[str, int]) -> Path:
    wc = P16.word_count(rendered_markdown())
    rel = lambda p: p.relative_to(ROOT).as_posix()  # noqa: E731
    write_json(MANIFEST, {
        "status": STATUS,
        "source_commit": SOURCE_COMMIT, "source_manuscript": SOURCE_MANUSCRIPT,
        "source_package": rel(P16_PACKAGE), "final_title": title(),
        "abstract_words": wc["abstract"], "main_text_words": wc["main_text"],
        "main_docx": rel(MAIN_DOCX), "supplementary_docx": rel(SUPP_DOCX), "supplementary_workbook": rel(SUPP_XLSX),
        "figures_dir": rel(FIGURES), "cover_letter_draft": rel(COVER_LETTER), "metadata_checklist": rel(METADATA),
        "external_information_request": rel(EXTERNAL_REQUEST),
        "reference_check": rel(REFERENCE_CHECK), "submission_checklist": rel(CHECKLIST),
        "external_open_items": N_METADATA_ITEMS, "manuscript_placeholders": len(P16.B.placeholders(P16.read_source())),
        "tables": N_TABLES, "figures": N_FIGURES, "supplementary_tables": N_SUPP_TABLES,
        "supplementary_figures": N_SUPP_FIGURES, "pages": pages,
        "qa_reports": [rel(p) for p in QA_REPORTS.values()],
        "sha256": {rel(p): sha256(p) for p in artifacts()}})
    return MANIFEST


def manifest_problems() -> list[str]:
    if not MANIFEST.is_file():
        return ["MANIFEST.json missing"]
    m = json.loads(MANIFEST.read_text(encoding="utf-8"))
    out = [] if m.get("status") == STATUS else [f"status {m.get('status')!r}"]
    if m.get("source_commit") != SOURCE_COMMIT or m.get("final_title") != title():
        out.append("source commit or title differs")
    for rel, digest in m.get("sha256", {}).items():
        p = ROOT / rel
        if not p.is_file():
            out.append(f"artifact missing: {rel}")
        elif sha256(p) != digest:
            out.append(f"hash mismatch: {rel}")
    listed = set(m.get("sha256", {}))
    out += [f"artifact not in manifest: {p.relative_to(ROOT).as_posix()}" for p in artifacts()
            if p.relative_to(ROOT).as_posix() not in listed]
    return out


# --- validation -----------------------------------------------------------------------------------------------------

FROZEN_SOURCES = ("paper/manuscript/manuscript_p16_final.md", "paper/manuscript/references.bib",
                  "paper/submission_p16", "paper/manuscript/COVER_LETTER_DRAFT.md")


def frozen_source_problems() -> list[str]:
    """The P16 manuscript, bibliography, package and P8-era cover letter are unchanged since the source commit."""
    import subprocess
    r = subprocess.run(["git", "diff", "--name-only", SOURCE_COMMIT, "--", *FROZEN_SOURCES], cwd=ROOT,
                       capture_output=True, text=True)
    if r.returncode:
        return [f"git diff failed: {r.stderr.strip()}"]
    return [f"changed since {SOURCE_COMMIT[:7]}: {n}" for n in r.stdout.split()]


def front_problems() -> list[str]:
    wc = P16.word_count(rendered_markdown())
    out = [] if wc["abstract"] == 198 else [f"abstract has {wc['abstract']} words (expected 198)"]
    if title() not in " ".join(docx_text(MAIN_DOCX).split()):
        out.append("final title missing from the main DOCX")
    return out


def validate() -> dict[str, list[str]]:
    for p in (MAIN_DOCX, SUPP_DOCX, SUPP_XLSX):
        if not p.is_file():
            return {"artifacts present": [f"missing: {p.relative_to(ROOT).as_posix()}"]}
    pages = json.loads(MANIFEST.read_text(encoding="utf-8")).get("pages", {}) if MANIFEST.is_file() else {}
    return {
        "frozen P16 sources unchanged since the source commit": frozen_source_problems(),
        "final title and 198-word abstract": front_problems(),
        "main DOCX structure (11 tables, 6 figures, 31 references, back matter, placeholders only, no Figure S5)":
            main_structure_problems(),
        "scientific equivalence with P16 (title, abstract, conclusions, every number)": equivalence_problems(),
        "supplementary DOCX and lossless workbook (S1-S44, Figures S1-S4)": supplementary_problems(),
        "Figures 1-6 files byte-identical to P16": figure_problems(),
        "privacy and document metadata": privacy_problems(MAIN_DOCX) + privacy_problems(SUPP_DOCX)
            + workbook_privacy_problems(),
        "cover letter (final title, no overstatement, placeholders, P8 draft unchanged)": cover_letter_problems(),
        "metadata checklist (18 open items covering the 16 placeholders)": metadata_problems(),
        "external information request (every item open and requested)": external_request_problems(),
        "QA reports (every rendered page recorded)": qa_problems(pages),
        "manifest (status, source commit, SHA-256 of every artifact)": manifest_problems(),
    }
